"""
services/knowledge_compiler/extractors/docx.py — extrator DOCX via
python-docx. Seções são delimitadas por parágrafos com estilo "Heading N";
parágrafos/listas viram conteúdo da seção corrente e tabelas são
serializadas como texto tabular simples, tudo na ordem real em que
aparecem no documento (python-docx só expõe `doc.paragraphs`/`doc.tables`
separadamente — por isso a iteração de baixo nível abaixo).
"""
import io

from docx import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from services.knowledge_compiler.extractors.base import DocumentExtractor
from services.knowledge_compiler.schemas import DocumentSection, ExtractedDocument

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _iter_block_items(doc):
    """Itera parágrafos e tabelas na ordem em que aparecem no corpo do
    documento (recipe padrão do python-docx para isso)."""
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


class DocxExtractor(DocumentExtractor):
    def supports(self, mime_type: str, filename: str) -> bool:
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        return ext == "docx" or mime_type == _DOCX_MIME

    def extract(self, conteudo: bytes) -> ExtractedDocument:
        doc = DocxDocument(io.BytesIO(conteudo))

        sections: list[DocumentSection] = []
        titulo_atual = "Documento"
        nivel_atual = 1
        buffer: list[str] = []
        ordem = 0

        def fechar_secao():
            nonlocal buffer, ordem
            texto = "\n".join(buffer).strip()
            if texto:
                sections.append(DocumentSection(
                    id=ordem, title=titulo_atual, content=texto, level=nivel_atual, order=ordem,
                ))
                ordem += 1
            buffer = []

        for item in _iter_block_items(doc):
            if isinstance(item, Table):
                linhas = [" | ".join(c.text.strip() for c in linha.cells) for linha in item.rows]
                if linhas:
                    buffer.append("\n".join(linhas))
                continue

            texto_para = item.text.strip()
            if not texto_para:
                continue

            estilo = (item.style.name or "").lower()
            if estilo.startswith("heading"):
                fechar_secao()
                titulo_atual = texto_para
                try:
                    nivel_atual = int(estilo.replace("heading", "").strip() or 1)
                except ValueError:
                    nivel_atual = 1
            else:
                prefixo = "- " if estilo.startswith("list") else ""
                buffer.append(prefixo + texto_para)

        fechar_secao()

        texto_completo = "\n\n".join(s.content for s in sections)
        return ExtractedDocument(text=texto_completo, metadata={}, sections=sections)
